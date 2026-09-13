from datetime import timedelta

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from openpyxl import load_workbook

from simple_crm.crm.lead_services import create_lead
from simple_crm.crm.models import Campaign, LeadPartyRole
from simple_crm.crm.services import create_organization_party
from simple_crm.identity.models import (
    IdentityProfile,
    Role,
    RoleAssignment,
    ScopeGrant,
    ScopeType,
)
from simple_crm.reporting.exports import (
    download_export,
    materialize_export,
    request_export,
)
from simple_crm.reporting.models import ExportAuditEvent, ExportState


def _actor(username: str, role_code: str) -> IdentityProfile:
    user = get_user_model().objects.create_user(username=username, password="safe-pass")
    profile = IdentityProfile.objects.get(user=user)
    RoleAssignment.objects.create(
        identity=profile, role=Role.objects.get(code=role_code)
    )
    campaign = Campaign.objects.get(code="GENERAL")
    ScopeGrant.objects.create(
        identity=profile, scope_type=ScopeType.CAMPAIGN, campaign=campaign
    )
    return profile


def _lead(actor: IdentityProfile, name: str, source: str = "web"):
    campaign = Campaign.objects.get(code="GENERAL")
    party, _ = create_organization_party(display_name=name)
    return create_lead(
        actor=actor,
        campaign=campaign,
        owner=actor,
        parties=((party, LeadPartyRole.ACCOUNT),),
        source_attribution={"source": source},
    )[0]


@pytest.mark.django_db
def test_export_requires_separate_permission_and_audits_denial() -> None:
    seller = _actor("export-seller", "sales_manager")
    with pytest.raises(PermissionDenied):
        request_export(
            actor=seller,
            definition={"and": [], "or": [], "sort": "created_desc"},
            fields=["lead_number"],
            export_format="CSV",
            idempotency_key="denied-1",
            reason="Revisión autorizada",
        )
    denied = ExportAuditEvent.objects.get(action="denied")
    assert denied.metadata == {"reason": "missing_export_permission"}


@pytest.mark.django_db
def test_export_snapshots_scope_materializes_csv_and_is_idempotent() -> None:
    actor = _actor("export-admin", "system_administrator")
    lead = _lead(actor, "=Cuenta peligrosa", source="=2+2")
    definition = {
        "and": [{"field": "stage", "operator": "eq", "value": "NUEVO"}],
        "or": [],
        "sort": "lead_number",
    }
    request = request_export(
        actor=actor,
        definition=definition,
        fields=["lead_number", "source", "created_date"],
        export_format="CSV",
        idempotency_key="export-idempotent",
        reason="Análisis comercial",
    )
    definition["and"] = []
    same = request_export(
        actor=actor,
        definition={
            "and": [{"field": "stage", "operator": "eq", "value": "NUEVO"}],
            "or": [],
            "sort": "lead_number",
        },
        fields=["lead_number", "source", "created_date"],
        export_format="CSV",
        idempotency_key="export-idempotent",
        reason="Análisis comercial",
    )
    assert same.pk == request.pk
    ready = materialize_export(actor=actor, export_id=request.pk)
    assert ready.state == ExportState.READY
    assert ready.row_count == 1
    assert ready.definition_snapshot["and"]
    assert ready.checksum
    response = download_export(actor=actor, export_id=request.pk)
    body = response.content.decode("utf-8")
    assert "'=2+2" in body
    assert "Cuenta peligrosa" not in body
    assert lead.lead_number in body
    assert ExportAuditEvent.objects.filter(
        export_request=request, action="downloaded"
    ).exists()


@pytest.mark.django_db
def test_xlsx_export_and_expiry_redact_no_sensitive_payload() -> None:
    actor = _actor("export-xlsx", "system_administrator")
    _lead(actor, "Cuenta XLSX", source="interna")
    request = request_export(
        actor=actor,
        definition={"and": [], "or": [], "sort": "created_desc"},
        fields=["lead_number", "campaign_code", "created_date"],
        export_format="XLSX",
        idempotency_key="export-xlsx",
        reason="Informe aprobado",
    )
    ready = materialize_export(actor=actor, export_id=request.pk)
    workbook = load_workbook(filename=__import__("io").BytesIO(ready.file_payload))
    assert workbook.active.max_row == 2
    assert "Correo" not in workbook.active[1][0].value
    type(request).objects.filter(pk=request.pk).update(
        expires_at=timezone.now() - timedelta(minutes=1)
    )
    with pytest.raises(PermissionDenied):
        download_export(actor=actor, export_id=request.pk)
    assert type(request).objects.get(pk=request.pk).state == ExportState.EXPIRED
